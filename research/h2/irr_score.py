"""Inter-rater agreement for the payer dictionary: percent agreement, Cohen's kappa, list of disagreements.

Usage: python -m research.h2.irr_score [rater_A.xlsx rater_B.xlsx]  (CSV files also accepted; default: round 2)
Target (proposal Table 9): 85% agreement or higher on canonical_payer.
Disagreements are also grouped by type, so the guide can be clarified where raters read a rule differently:
"plan names insurer" = one rater chose the insurer named only in the plan column and the other chose "other".
"""
import csv
import re
import sys
from collections import Counter

from pipeline.payers import norm

PLAN_INSURER = ((r"\bcigna\b", "cigna"), (r"\baetna\b", "aetna"),
                (r"\bbcbs\b|\bblue cross\b|\bblue card\b", "blue cross blue shield of arizona"),
                (r"\buhc\b|\bunited ?healthcare\b", "unitedhealthcare"), (r"\bhumana\b", "humana"))


def kind(row: dict, a: str, b: str) -> str:
    """Type of a canonical_payer disagreement between labels a and b (both casefolded)."""
    plan = norm(row["raw_plan"])
    named = {canon for rx, canon in PLAN_INSURER if re.search(rx, plan)}
    if {a, b} & named and "other" in (a, b):
        return "plan names insurer"
    if "banner health plans" in (a, b) and "banner" in row["hospital"].casefold() \
            and "banner" not in norm(row["raw_payer"] + " " + row["raw_plan"]):
        return "hospital name used as payer"
    return "other disagreement"


def load(path: str) -> list[dict]:
    """Read a rater file: the Excel copy (.xlsx, sheet "label") or the CSV."""
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        rows = list(load_workbook(path, read_only=True)["label"].iter_rows(values_only=True))
        head = [str(h) for h in rows[0]]
        return [{k: ("" if v is None else str(v)) for k, v in zip(head, r, strict=False)} for r in rows[1:] if any(r)]
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def kappa(a: list[str], b: list[str]) -> float:
    n = len(a)
    po = sum(x == y for x, y in zip(a, b, strict=True)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def main() -> None:
    pa, pb = (sys.argv[1:3] if len(sys.argv) >= 3 else
              ("research/h2/irr_round2_rater_A.xlsx", "research/h2/irr_round2_rater_B.xlsx"))
    A, B = load(pa), load(pb)
    assert [(r["raw_payer"], r["raw_plan"]) for r in A] == [(r["raw_payer"], r["raw_plan"]) for r in B], \
        "the two files must list the same strings in the same order"
    empty = [i for i, (x, y) in enumerate(zip(A, B, strict=True))
             if not x["canonical_payer"].strip() or not y["canonical_payer"].strip()]
    if empty:
        sys.exit(f"{len(empty)} rows are not labelled yet (first: row {empty[0] + 2})")
    for field in ("canonical_payer", "plan_type"):
        a = [r[field].strip().casefold() for r in A]
        b = [r[field].strip().casefold() for r in B]
        agree = sum(x == y for x, y in zip(a, b, strict=True))
        print(f"{field}: {agree}/{len(a)} = {agree / len(a):.1%} agreement, Cohen's kappa = {kappa(a, b):.3f}")
    print("target: 85% agreement on canonical_payer")
    kinds = Counter(kind(x, x["canonical_payer"].strip().casefold(), y["canonical_payer"].strip().casefold())
                    for x, y in zip(A, B, strict=True)
                    if x["canonical_payer"].strip().casefold() != y["canonical_payer"].strip().casefold())
    print("canonical_payer disagreements by type:", dict(kinds.most_common()))
    print("\nDisagreements (canonical_payer or plan_type):")
    for x, y in zip(A, B, strict=True):
        if (x["canonical_payer"].strip().casefold() != y["canonical_payer"].strip().casefold()
                or x["plan_type"].strip().casefold() != y["plan_type"].strip().casefold()):
            print(f"- {x['raw_payer']} | {x['raw_plan']}: A={x['canonical_payer']}/{x['plan_type']}  "
                  f"B={y['canonical_payer']}/{y['plan_type']}")


if __name__ == "__main__":
    main()
