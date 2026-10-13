"""Draw the 200 payer strings for the inter-rater test. Seed and method are fixed here and in REPORT.md.

Method: from the distinct (raw_payer, raw_plan) strings in data/processed/payer_aliases.csv, draw 100 without
replacement with probability proportional to row_count, then 100 uniformly from the strings not yet drawn
(so rare strings are included). Order is shuffled so raters cannot tell which half a string came from.
"""
import csv
import random

SEED = 20261012
N_WEIGHTED, N_UNIFORM = 100, 100

with open("data/processed/payer_aliases.csv", newline="") as f:
    pool = list(csv.DictReader(f))
with open("data/processed/payer_strings.csv", newline="") as f:
    where: dict = {}
    for r in csv.DictReader(f):
        where.setdefault((r["raw_payer"], r["raw_plan"]), []).append((int(r["row_count"]), r["hospital"]))

rng = random.Random(SEED)
remaining = list(pool)
chosen = []
for _ in range(min(N_WEIGHTED, len(remaining))):
    pick = rng.choices(remaining, weights=[int(r["row_count"]) for r in remaining])[0]
    chosen.append(pick)
    remaining.remove(pick)
chosen += rng.sample(remaining, min(N_UNIFORM, len(remaining)))
rng.shuffle(chosen)

PLAN_TYPES = ["commercial", "medicare advantage", "medicare supplement", "ahcccs", "marketplace", "government",
              "workers comp", "rental network", "unknown"]


def write_xlsx(name: str) -> None:
    """Excel copy with drop-down lists, so a rater only picks values (same rows as the CSV)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from pipeline.payers import CANONICAL
    wb = Workbook()
    ws = wb.active
    ws.title = "label"
    lists = wb.create_sheet("lists")
    for i, v in enumerate(CANONICAL, 1):
        lists.cell(i, 1, v)
    for i, v in enumerate(PLAN_TYPES, 1):
        lists.cell(i, 2, v)
    ws.append(["raw_payer", "raw_plan", "hospital", "canonical_payer", "plan_type", "notes"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in chosen:
        ws.append([r["raw_payer"], r["raw_plan"], max(where[(r["raw_payer"], r["raw_plan"])])[1], "", "", ""])
    n = len(chosen) + 1
    dv1 = DataValidation(type="list", formula1=f"=lists!$A$1:$A${len(CANONICAL)}", allow_blank=False)
    dv2 = DataValidation(type="list", formula1=f"=lists!$B$1:$B${len(PLAN_TYPES)}", allow_blank=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"D2:D{n}")
    dv2.add(f"E2:E{n}")
    fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        for c in row:
            c.fill = fill
    for col, w in zip("ABCDEF", (38, 44, 34, 40, 20, 30), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(f"research/h2/{name}")


for name in ("irr_rater_A.xlsx", "irr_rater_B.xlsx"):
    write_xlsx(name)

for name in ("irr_rater_A.csv", "irr_rater_B.csv"):
    with open(f"research/h2/{name}", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["raw_payer", "raw_plan", "hospital", "canonical_payer", "plan_type", "notes"])
        for r in chosen:
            hosp = max(where[(r["raw_payer"], r["raw_plan"])])[1]
            w.writerow([r["raw_payer"], r["raw_plan"], hosp, "", "", ""])
print(f"{len(chosen)} strings written (pool {len(pool)}), seed {SEED}")
